"""Planilhas (XLSX e CSV).

Planilha não é prosa: cada linha é um registro. O parser gera um bloco por
janela de linhas, sempre repetindo o cabeçalho, para que um chunk isolado
continue interpretável -- sem isso, o LLM recebe números sem saber a que
coluna pertencem.

Cada aba vira uma `section`, o que permite citar "planilha X, aba Y".
"""

from __future__ import annotations

import csv
from pathlib import Path

from enterprise_rag.ingestion.parsers.base import DocumentParser, ParserError
from enterprise_rag.models.document import ParsedBlock

ROWS_PER_BLOCK = 25


class TabularParser(DocumentParser):
    name = "tabular"
    extensions = (".xlsx", ".xlsm", ".csv")

    def parse(self, path: Path) -> list[ParsedBlock]:
        sheets = {"csv": _read_csv(path)} if path.suffix.lower() == ".csv" else _read_xlsx(path)

        blocks: list[ParsedBlock] = []
        for sheet_name, rows in sheets.items():
            nonempty_rows = [row for row in rows if any(str(cell).strip() for cell in row)]
            if not nonempty_rows:
                continue
            header = [str(cell).strip() for cell in nonempty_rows[0]]
            for start in range(1, len(nonempty_rows), ROWS_PER_BLOCK):
                window = nonempty_rows[start : start + ROWS_PER_BLOCK]
                lines = ["| " + " | ".join(header) + " |"]
                lines.append("| " + " | ".join("---" for _ in header) + " |")
                for row in window:
                    cells = [str(cell).strip() if cell is not None else "" for cell in row]
                    cells += [""] * (len(header) - len(cells))
                    lines.append("| " + " | ".join(cells[: len(header)]) + " |")
                blocks.append(
                    ParsedBlock(
                        text="\n".join(lines),
                        section=sheet_name,
                        kind="table",
                        metadata={"first_row": start, "rows": len(window)},
                    )
                )
        if not blocks:
            raise ParserError(f"nenhuma linha util em {path.name}")
        return blocks


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8", errors="replace", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(handle, dialect)]


def _read_xlsx(path: Path) -> dict[str, list[list[object]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ParserError("openpyxl nao instalado; rode: pip install '.[parsers]'") from exc

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        return {
            str(sheet.title): [list(row) for row in sheet.iter_rows(values_only=True)]
            for sheet in workbook.worksheets
        }
    finally:
        workbook.close()
