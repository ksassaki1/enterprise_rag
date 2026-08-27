"""DOCX via python-docx.

Word já entrega a estrutura pronta: estilos `Heading N` viram `section` e
tabelas viram blocos `kind="table"`, serializados em Markdown para que o LLM
consiga ler linha e coluna. Achatar tabela em texto corrido é uma das formas
mais rápidas de destruir a resposta de uma pergunta numérica.
"""

from __future__ import annotations

from pathlib import Path

from enterprise_rag.ingestion.parsers.base import DocumentParser, ParserError
from enterprise_rag.models.document import ParsedBlock


class DocxParser(DocumentParser):
    name = "docx"
    extensions = (".docx",)

    def parse(self, path: Path) -> list[ParsedBlock]:
        try:
            import docx
        except ImportError as exc:
            raise ParserError(
                "python-docx nao instalado; rode: pip install '.[parsers]'"
            ) from exc

        document = docx.Document(str(path))
        blocks: list[ParsedBlock] = []
        section: str | None = None
        buffer: list[str] = []

        def flush() -> None:
            text = "\n".join(buffer).strip()
            buffer.clear()
            if text:
                blocks.append(ParsedBlock(text=text, section=section, kind="paragraph"))

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            if not text:
                flush()
                continue
            if style.startswith("heading") or style in {"title", "subtitle"}:
                flush()
                section = text
                blocks.append(ParsedBlock(text=text, section=section, kind="heading"))
            else:
                buffer.append(text)
        flush()

        for index, table in enumerate(document.tables, start=1):
            markdown = _table_to_markdown(table)
            if markdown:
                blocks.append(
                    ParsedBlock(
                        text=markdown,
                        section=section,
                        kind="table",
                        metadata={"table_index": index},
                    )
                )

        if not blocks:
            raise ParserError(f"nenhum texto extraido de {path.name}")
        return blocks


def _table_to_markdown(table: object) -> str:
    rows = [
        [cell.text.strip().replace("\n", " ") for cell in row.cells]
        for row in table.rows  # type: ignore[attr-defined]
    ]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |"]
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)
